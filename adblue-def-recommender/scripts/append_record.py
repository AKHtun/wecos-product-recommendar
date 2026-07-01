"""Append a Recommendation Record row to the working log.

Reads a JSON payload from stdin OR a file path, generates the next Record ID
(REC-YYYYMMDD-NNN), copies the template if the log doesn't exist yet, and
appends the row.

Usage:
    python3 scripts/append_record.py path/to/payload.json
    cat payload.json | python3 scripts/append_record.py

Payload fields (all optional except where noted; missing fields written blank):
{
  "dsr":                   "Jane Doe",
  "customer":              "Acme Mining",
  "incumbent_brand":       "Shell",
  "incumbent_product":     "Tellus S2 M 46",       # REQUIRED
  "application":           "Hydraulic",            # REQUIRED
  "equipment_oem_model":   "Caterpillar 390F",
  "operating_temp":        "-10 to +80 C",
  "viscosity_nlgi":        "ISO VG 46",
  "base_oil_type":         "Mineral (Group II)",
  "industry_standards":    "DIN 51524-2 HLP; ISO 11158 HM",
  "oem_approvals":         "Denison HF-0; Eaton 35VQ25",
  "suppliers_searched":    "ExxonMobil, Castrol, TotalEnergies",
  "recommended_1":         "ExxonMobil Mobil DTE 25 Ultra",
  "recommended_2":         "Castrol Hyspin AWS 46",
  "recommended_3":         "TotalEnergies Azolla ZS 46",
  "match_score":           6,
  "special_case_warning":  "",
  "pds_source_urls":       "https://...; https://...",
  "information_gaps":      "",
  "report_pdf_path":       "output/recommendations/REC-20260614-001.pdf",
  "dle_notes":             "All three meet Denison HF-0; rank by local availability."
}
"""
from __future__ import annotations
import json
import sys
import shutil
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "templates" / "recommendation_record.xlsx"
OUT_DIR = Path("output")
LOG = OUT_DIR / "recommendation_record.xlsx"

COLUMNS = [
    "record_id", "date", "dsr", "customer",
    "incumbent_brand", "incumbent_product", "application",
    "equipment_oem_model", "operating_temp", "viscosity_nlgi",
    "base_oil_type", "industry_standards", "oem_approvals",
    "suppliers_searched", "recommended_1", "recommended_2",
    "recommended_3", "match_score", "special_case_warning",
    "pds_source_urls", "information_gaps", "report_pdf_path",
    "dle_notes",
]

REQUIRED = {"incumbent_product", "application"}


def load_payload() -> dict:
    if len(sys.argv) >= 2:
        return json.loads(Path(sys.argv[1]).read_text())
    return json.loads(sys.stdin.read())


def ensure_log() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if not LOG.exists():
        shutil.copy(TEMPLATE, LOG)
    return LOG


def next_record_id(ws) -> str:
    today = datetime.now().strftime("%Y%m%d")
    prefix = f"REC-{today}-"
    used = []
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        rid = row[0]
        if isinstance(rid, str) and rid.startswith(prefix):
            try:
                used.append(int(rid.rsplit("-", 1)[1]))
            except ValueError:
                pass
    nxt = max(used) + 1 if used else 1
    return f"{prefix}{nxt:03d}"


def first_empty_row(ws) -> int:
    # Row 2 of the fresh template is a styled blank placeholder; reuse it.
    for r in range(2, ws.max_row + 2):
        if all(ws.cell(row=r, column=c).value in (None, "") for c in range(1, len(COLUMNS) + 1)):
            return r
    return ws.max_row + 1


def main() -> None:
    payload = load_payload()
    missing = REQUIRED - {k for k, v in payload.items() if v}
    if missing:
        sys.exit(f"Missing required fields: {sorted(missing)}")

    log = ensure_log()
    wb = load_workbook(log)
    ws = wb.active

    rid = next_record_id(ws)
    payload["record_id"] = rid
    payload.setdefault("date", datetime.now().strftime("%Y-%m-%d"))

    row = first_empty_row(ws)
    body_font = Font(name="Arial", size=10)
    align = Alignment(vertical="top", wrap_text=True)
    thin = Side(border_style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, key in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=row, column=col_idx, value=payload.get(key, ""))
        cell.font = body_font
        cell.alignment = align
        cell.border = border

    wb.save(log)
    print(json.dumps({"record_id": rid, "row": row, "log": str(log)}))


if __name__ == "__main__":
    main()
