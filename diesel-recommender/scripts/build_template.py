"""Build the Recommendation Record Excel template.

Run once to produce templates/recommendation_record.xlsx (header row + formatting only).
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path

HEADERS = [
    ("Record ID",            14),
    ("Date",                 12),
    ("DSR",                  16),
    ("Customer / End User",  22),
    ("Incumbent Brand",      16),
    ("Incumbent Product",    26),
    ("Application",          18),
    ("Equipment OEM/Model",  22),
    ("Operating Temp",       14),
    ("Viscosity / NLGI",     16),
    ("Base Oil Type",        14),
    ("Industry Standards",   24),
    ("OEM Approvals",        24),
    ("Suppliers Searched",   24),
    ("Recommended #1",       28),
    ("Recommended #2",       28),
    ("Recommended #3",       28),
    ("Match Score (0-6)",    14),
    ("Special-Case Warning", 28),
    ("PDS Source URLs",      40),
    ("Information Gaps",     28),
    ("Report PDF Path",      36),
    ("DLE Notes",            36),
]

OUT = Path(__file__).resolve().parent.parent / "templates" / "recommendation_record.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "Recommendation Record"

header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", start_color="1F4E78")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin = Side(border_style="thin", color="999999")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for col, (label, width) in enumerate(HEADERS, start=1):
    cell = ws.cell(row=1, column=col, value=label)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = border
    ws.column_dimensions[get_column_letter(col)].width = width

ws.row_dimensions[1].height = 36
ws.freeze_panes = "A2"

# Sample blank row with default body font
body_font = Font(name="Arial", size=10)
for col in range(1, len(HEADERS) + 1):
    cell = ws.cell(row=2, column=col, value=None)
    cell.font = body_font
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    cell.border = border

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"Wrote {OUT}")
